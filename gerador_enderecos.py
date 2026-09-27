from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def gerar_enderecos():
	enderecos = []

	for numero_buffer in range(1, 6):
		for numero_rua in range(1, 21):
			for numero_canalizacao in range(2, 11, 2):
				buffer = f"B{numero_buffer:02d}"
				rua = f"R{numero_rua:02d}"
				canalizacao = f"BA{numero_canalizacao:02d}"
				codigo = f"{buffer}-{rua}-{canalizacao}"

				enderecos.append(
					{
						"codigo": codigo,
						"buffer": buffer,
						"rua": rua,
						"canalizacao": canalizacao,
						"status": "Disponível",
						"produto": "",
						"quantidade": 0,
					}
				)

	return enderecos


def criar_planilha(enderecos):
	planilha = Workbook()
	pagina = planilha.active
	pagina.title = "Endereços"
	pagina.append(
		["Código", "Buffer", "Rua", "Canalização", "Status", "Produto", "Quantidade"]
	)

	for endereco in enderecos:
		pagina.append(
			[
				endereco["codigo"],
				endereco["buffer"],
				endereco["rua"],
				endereco["canalizacao"],
				endereco["status"],
				endereco["produto"],
				endereco["quantidade"],
			]
		)

	preenchimento = PatternFill(fill_type="solid", fgColor="1F4E78")
	fonte = Font(color="FFFFFF", bold=True)
	for celula in pagina[1]:
		celula.fill = preenchimento
		celula.font = fonte
		celula.alignment = Alignment(horizontal="center")

	for coluna, largura in {
		"A": 22,
		"B": 12,
		"C": 12,
		"D": 16,
		"E": 15,
		"F": 30,
		"G": 14,
	}.items():
		pagina.column_dimensions[coluna].width = largura

	pagina.freeze_panes = "A2"
	pagina.auto_filter.ref = pagina.dimensions
	nome_arquivo = "enderecos_logisticos.xlsx"
	planilha.save(nome_arquivo)
	return nome_arquivo


if __name__ == "__main__":
	lista_enderecos = gerar_enderecos()
	arquivo_criado = criar_planilha(lista_enderecos)
	print("Planilha criada com sucesso!")
	print(f"Arquivo: {arquivo_criado}")
	print(f"Total de endereços: {len(lista_enderecos)}")
