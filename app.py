from pathlib import Path
from io import BytesIO
from datetime import datetime

import pandas as pd
import streamlit as st
from openpyxl import Workbook, load_workbook

from gerador_enderecos import criar_planilha, gerar_enderecos


ARQUIVO_EXCEL = Path("enderecos_logisticos.xlsx")
ARQUIVO_HISTORICO = Path("historico_movimentacoes.xlsx")

st.set_page_config(
    page_title="Sistema de Endereçamento Logístico",
    page_icon="📦",
    layout="wide",
)


def criar_arquivo_se_nao_existir():
    if not ARQUIVO_EXCEL.exists():
        enderecos_iniciais = gerar_enderecos()
        criar_planilha(enderecos_iniciais)


def carregar_enderecos():
    criar_arquivo_se_nao_existir()

    tabela = pd.read_excel(
        ARQUIVO_EXCEL,
        engine="openpyxl",
    )

    tabela.columns = [
        "codigo",
        "buffer",
        "rua",
        "canalizacao",
        "status",
        "produto",
        "quantidade",
    ]

    tabela["produto"] = tabela["produto"].fillna("")
    tabela["quantidade"] = tabela["quantidade"].fillna(0)

    return tabela


def salvar_enderecos(tabela):
    tabela_para_salvar = tabela.copy()

    tabela_para_salvar.columns = [
        "Código",
        "Buffer",
        "Rua",
        "Canalização",
        "Status",
        "Produto",
        "Quantidade",
    ]

    tabela_para_salvar.to_excel(
        ARQUIVO_EXCEL,
        index=False,
        engine="openpyxl",
    )

def registrar_historico(
    tipo,
    codigo,
    produto,
    quantidade_movimentada,
    quantidade_anterior,
    quantidade_final,
    observacao="",
):
    if ARQUIVO_HISTORICO.exists():
        planilha = load_workbook(ARQUIVO_HISTORICO)
        pagina = planilha.active
    else:
        planilha = Workbook()
        pagina = planilha.active
        pagina.title = "Histórico"

        pagina.append(
            [
                "Data e hora",
                "Tipo",
                "Endereço",
                "Produto",
                "Quantidade movimentada",
                "Quantidade anterior",
                "Quantidade final",
                "Observação",
            ]
        )

    pagina.append(
        [
            datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
            tipo,
            codigo,
            produto,
            quantidade_movimentada,
            quantidade_anterior,
            quantidade_final,
            observacao,
        ]
    )

    planilha.save(ARQUIVO_HISTORICO)

def carregar_historico():
    colunas = [
        "Data e hora",
        "Tipo",
        "Endereço",
        "Produto",
        "Quantidade movimentada",
        "Quantidade anterior",
        "Quantidade final",
        "Observação",
    ]

    if not ARQUIVO_HISTORICO.exists():
        return pd.DataFrame(columns=colunas)

    historico = pd.read_excel(
        ARQUIVO_HISTORICO,
        engine="openpyxl",
    )

    for coluna in colunas:
        if coluna not in historico.columns:
            historico[coluna] = ""

    historico = historico[colunas]

    historico["Data e hora"] = pd.to_datetime(
        historico["Data e hora"],
        dayfirst=True,
        errors="coerce",
    )

    historico["Produto"] = historico[
        "Produto"
    ].fillna("")

    historico["Observação"] = historico[
        "Observação"
    ].fillna("")

    return historico

def criar_excel_para_download(tabela, nome_pagina):
    arquivo_memoria = BytesIO()

    with pd.ExcelWriter(
        arquivo_memoria,
        engine="openpyxl",
    ) as escritor:
        tabela.to_excel(
            escritor,
            index=False,
            sheet_name=nome_pagina,
        )

        pagina = escritor.sheets[nome_pagina]

        pagina.freeze_panes = "A2"
        pagina.auto_filter.ref = pagina.dimensions

        for coluna in pagina.columns:
            maior_tamanho = 0
            letra_coluna = coluna[0].column_letter

            for celula in coluna:
                valor = "" if celula.value is None else str(
                    celula.value
                )

                maior_tamanho = max(
                    maior_tamanho,
                    len(valor),
                )

            largura = min(maior_tamanho + 3, 40)

            pagina.column_dimensions[
                letra_coluna
            ].width = largura

    arquivo_memoria.seek(0)

    return arquivo_memoria.getvalue()


enderecos = carregar_enderecos()


st.title("📦 Sistema de Endereçamento Logístico")

st.write(
    "Controle de buffers, ruas, canalizações, "
    "produtos e posições de armazenagem."
)

st.divider()


total_enderecos = len(enderecos)

total_disponiveis = len(
    enderecos[enderecos["status"] == "Disponível"]
)

total_ocupados = len(
    enderecos[enderecos["status"] == "Ocupado"]
)

if total_enderecos > 0:
    taxa_ocupacao = total_ocupados / total_enderecos * 100
else:
    taxa_ocupacao = 0


coluna1, coluna2, coluna3, coluna4 = st.columns(4)

coluna1.metric(
    "Total de endereços",
    total_enderecos,
)

coluna2.metric(
    "Endereços disponíveis",
    total_disponiveis,
)

coluna3.metric(
    "Endereços ocupados",
    total_ocupados,
)

coluna4.metric(
    "Taxa de ocupação",
    f"{taxa_ocupacao:.1f}%",
)


st.divider()


(
    pagina_cadastro,
    pagina_consulta,
    pagina_movimentacao,
    pagina_liberacao,
    pagina_historico,
    pagina_dashboard,
    pagina_relatorios,
) = st.tabs(
    [
        "Cadastrar produto",
        "Consultar endereços",
        "Movimentar estoque",
        "Liberar endereço",
        "Histórico",
        "Dashboard",
        "Relatórios",
    ]
)


with pagina_cadastro:
    st.subheader("Cadastrar produto em um endereço")

    enderecos_disponiveis = enderecos[
        enderecos["status"] == "Disponível"
    ]

    if enderecos_disponiveis.empty:
        st.warning(
            "Não existem endereços disponíveis para cadastro."
        )

    else:
        lista_codigos = (
            enderecos_disponiveis["codigo"]
            .sort_values()
            .tolist()
        )

        with st.form("formulario_cadastro"):
            codigo_selecionado = st.selectbox(
                "Endereço disponível",
                lista_codigos,
            )

            nome_produto = st.text_input(
                "Nome do produto",
                placeholder="Exemplo: Caixa de eletrônicos",
            )

            quantidade = st.number_input(
                "Quantidade",
                min_value=1,
                step=1,
            )

            botao_cadastrar = st.form_submit_button(
                "Cadastrar produto",
                type="primary",
            )

        if botao_cadastrar:
            produto_limpo = nome_produto.strip()

            if not produto_limpo:
                st.error(
                    "Digite o nome do produto antes de cadastrar."
                )

            else:
                endereco_encontrado = (
                    enderecos["codigo"] == codigo_selecionado
                )

                enderecos.loc[
                    endereco_encontrado,
                    "produto",
                ] = produto_limpo

                enderecos.loc[
                    endereco_encontrado,
                    "quantidade",
                ] = int(quantidade)

                enderecos.loc[
                    endereco_encontrado,
                    "status",
                ] = "Ocupado"

                salvar_enderecos(enderecos)
                registrar_historico(
                    tipo="Cadastro",
                    codigo=codigo_selecionado,
                    produto=produto_limpo,
                    quantidade_movimentada=int(quantidade),
                    quantidade_anterior=0,
                    quantidade_final=int(quantidade),
                    observacao="Cadastro inicial do produto",
                )

                st.write("Histórico registrado com sucesso.")
                
                st.success(
                    f"Produto cadastrado no endereço "
                    f"{codigo_selecionado}!"
                )

                st.rerun()


with pagina_consulta:
    st.subheader("Consultar endereços")

    coluna_filtro1, coluna_filtro2, coluna_filtro3 = (
        st.columns(3)
    )

    with coluna_filtro1:
        buffers = ["Todos"] + sorted(
            enderecos["buffer"].unique().tolist()
        )

        buffer_selecionado = st.selectbox(
            "Buffer",
            buffers,
        )

    with coluna_filtro2:
        ruas = ["Todas"] + sorted(
            enderecos["rua"].unique().tolist()
        )

        rua_selecionada = st.selectbox(
            "Rua",
            ruas,
        )

    with coluna_filtro3:
        status_selecionado = st.selectbox(
            "Status",
            [
                "Todos",
                "Disponível",
                "Ocupado",
            ],
        )

    pesquisa = st.text_input(
        "Pesquisar por código ou produto",
        placeholder="Exemplo: B01-R01-BA02",
    )

    enderecos_filtrados = enderecos.copy()

    if buffer_selecionado != "Todos":
        enderecos_filtrados = enderecos_filtrados[
            enderecos_filtrados["buffer"]
            == buffer_selecionado
        ]

    if rua_selecionada != "Todas":
        enderecos_filtrados = enderecos_filtrados[
            enderecos_filtrados["rua"]
            == rua_selecionada
        ]

    if status_selecionado != "Todos":
        enderecos_filtrados = enderecos_filtrados[
            enderecos_filtrados["status"]
            == status_selecionado
        ]

    if pesquisa:
        pesquisa_normalizada = pesquisa.strip()

        filtro_codigo = (
            enderecos_filtrados["codigo"]
            .str.contains(
                pesquisa_normalizada,
                case=False,
                na=False,
            )
        )

        filtro_produto = (
            enderecos_filtrados["produto"]
            .str.contains(
                pesquisa_normalizada,
                case=False,
                na=False,
            )
        )

        enderecos_filtrados = enderecos_filtrados[
            filtro_codigo | filtro_produto
        ]

    st.write(
        f"**Endereços encontrados:** "
        f"{len(enderecos_filtrados)}"
    )

    tabela_exibicao = enderecos_filtrados.rename(
        columns={
            "codigo": "Código",
            "buffer": "Buffer",
            "rua": "Rua",
            "canalizacao": "Canalização",
            "status": "Status",
            "produto": "Produto",
            "quantidade": "Quantidade",
        }
    )

    st.dataframe(
        tabela_exibicao,
        width="stretch",
        hide_index=True,
    )

with pagina_movimentacao:
    st.subheader("Movimentar estoque")

    st.write(
        "Registre a entrada ou a saída de unidades "
        "em um endereço ocupado."
    )

    enderecos_ocupados_movimentacao = enderecos[
        enderecos["status"] == "Ocupado"
    ].copy()

    if enderecos_ocupados_movimentacao.empty:
        st.info(
            "Não existem endereços ocupados para movimentar."
        )

    else:
        enderecos_ocupados_movimentacao["descricao"] = (
            enderecos_ocupados_movimentacao["codigo"]
            + " | "
            + enderecos_ocupados_movimentacao["produto"]
            + " | Quantidade: "
            + enderecos_ocupados_movimentacao["quantidade"]
            .astype(int)
            .astype(str)
        )

        lista_movimentacao = (
            enderecos_ocupados_movimentacao["descricao"]
            .sort_values()
            .tolist()
        )

        endereco_movimentacao = st.selectbox(
            "Selecione o endereço ocupado",
            lista_movimentacao,
            key="endereco_movimentacao",
        )

        codigo_movimentacao = endereco_movimentacao.split(
            " | "
        )[0]

        linha_movimentacao = enderecos[
            enderecos["codigo"] == codigo_movimentacao
        ].iloc[0]

        quantidade_atual = int(
            linha_movimentacao["quantidade"]
        )

        st.write("### Situação atual")

        coluna_endereco, coluna_produto, coluna_estoque = (
            st.columns(3)
        )

        coluna_endereco.metric(
            "Endereço",
            linha_movimentacao["codigo"],
        )

        coluna_produto.metric(
            "Produto",
            linha_movimentacao["produto"],
        )

        coluna_estoque.metric(
            "Quantidade atual",
            quantidade_atual,
        )

        with st.form("formulario_movimentacao"):
            tipo_movimentacao = st.radio(
                "Tipo de movimentação",
                [
                    "Entrada",
                    "Saída",
                ],
                horizontal=True,
            )

            quantidade_movimentada = st.number_input(
                "Quantidade da movimentação",
                min_value=1,
                step=1,
            )

            observacao_movimentacao = st.text_input(
                "Observação",
                placeholder=(
                    "Exemplo: recebimento, separação "
                    "ou ajuste de estoque"
                ),
            )

            confirmar_movimentacao = st.checkbox(
                "Confirmo os dados desta movimentação."
            )

            botao_movimentar = st.form_submit_button(
                "Registrar movimentação",
                type="primary",
            )

        if botao_movimentar:
            quantidade_informada = int(
                quantidade_movimentada
            )

            if not confirmar_movimentacao:
                st.warning(
                    "Marque a caixa de confirmação antes "
                    "de registrar a movimentação."
                )

            elif (
                tipo_movimentacao == "Saída"
                and quantidade_informada > quantidade_atual
            ):
                st.error(
                    "A quantidade de saída é maior que "
                    "a quantidade disponível no endereço."
                )

            else:
                if tipo_movimentacao == "Entrada":
                    nova_quantidade = (
                        quantidade_atual
                        + quantidade_informada
                    )

                else:
                    nova_quantidade = (
                        quantidade_atual
                        - quantidade_informada
                    )

                endereco_encontrado = (
                    enderecos["codigo"]
                    == codigo_movimentacao
                )

                enderecos.loc[
                    endereco_encontrado,
                    "quantidade",
                ] = nova_quantidade

                if nova_quantidade == 0:
                    enderecos.loc[
                        endereco_encontrado,
                        "produto",
                    ] = ""

                    enderecos.loc[
                        endereco_encontrado,
                        "status",
                    ] = "Disponível"

                try:
                    salvar_enderecos(enderecos)
                    registrar_historico(
                        tipo=tipo_movimentacao,
                        codigo=codigo_movimentacao,
                        produto=linha_movimentacao["produto"],
                        quantidade_movimentada=quantidade_informada,
                        quantidade_anterior=quantidade_atual,
                        quantidade_final=nova_quantidade,
                        observacao=observacao_movimentacao.strip(),
                    )

                except PermissionError:
                    st.error(
                        "Não foi possível salvar. Feche a "
                        "planilha enderecos_logisticos.xlsx "
                        "e tente novamente."
                    )

                else:
                    st.success(
                        f"{tipo_movimentacao} registrada "
                        f"com sucesso no endereço "
                        f"{codigo_movimentacao}."
                    )

                    if nova_quantidade == 0:
                        st.info(
                            "A quantidade chegou a zero. "
                            "O endereço foi liberado "
                            "automaticamente."
                        )

                    else:
                        st.write(
                            f"**Nova quantidade:** "
                            f"{nova_quantidade}"
                        )

                    st.rerun()

with pagina_liberacao:  
    st.subheader("Liberar um endereço ocupado")

    st.write(
        "Selecione um endereço ocupado para retirar o produto "
        "e tornar a posição disponível novamente."
    )

    enderecos_ocupados = enderecos[
        enderecos["status"] == "Ocupado"
    ].copy()

    if enderecos_ocupados.empty:
        st.info(
            "Não existem endereços ocupados para liberar."
        )

    else:
        enderecos_ocupados["descricao"] = (
            enderecos_ocupados["codigo"]
            + " | "
            + enderecos_ocupados["produto"]
            + " | Quantidade: "
            + enderecos_ocupados["quantidade"]
            .astype(int)
            .astype(str)
        )

        lista_enderecos_ocupados = (
            enderecos_ocupados["descricao"]
            .sort_values()
            .tolist()
        )

        endereco_para_liberar = st.selectbox(
            "Selecione o endereço ocupado",
            lista_enderecos_ocupados,
        )

        codigo_para_liberar = endereco_para_liberar.split(
            " | "
        )[0]

        registro_selecionado = enderecos[
            enderecos["codigo"] == codigo_para_liberar
        ].iloc[0]

        st.write("### Informações da posição")

        coluna_endereco, coluna_produto, coluna_quantidade = (
            st.columns(3)
        )

        coluna_endereco.metric(
            "Endereço",
            registro_selecionado["codigo"],
        )

        coluna_produto.metric(
            "Produto",
            registro_selecionado["produto"],
        )

        coluna_quantidade.metric(
            "Quantidade",
            int(registro_selecionado["quantidade"]),
        )

        confirmar_liberacao = st.checkbox(
            "Confirmo que desejo retirar o produto "
            "e liberar este endereço."
        )

        botao_liberar = st.button(
            "Liberar endereço",
            type="primary",
        )

        if botao_liberar:
            if not confirmar_liberacao:
                st.warning(
                    "Marque a caixa de confirmação antes "
                    "de liberar o endereço."
                )

            else:
                endereco_encontrado = (
                    enderecos["codigo"] == codigo_para_liberar
                )

                enderecos.loc[
                    endereco_encontrado,
                    "produto",
                ] = ""

                enderecos.loc[
                    endereco_encontrado,
                    "quantidade",
                ] = 0

                enderecos.loc[
                    endereco_encontrado,
                    "status",
                ] = "Disponível"

                try:
                    salvar_enderecos(enderecos)

                    registrar_historico(
                        tipo="Liberação",
                        codigo=codigo_para_liberar,
                        produto=registro_selecionado["produto"],
                        quantidade_movimentada=int(
                            registro_selecionado["quantidade"]
                        ),
                        quantidade_anterior=int(
                            registro_selecionado["quantidade"]
                        ),
                        quantidade_final=0,
                        observacao="Liberação manual do endereço",
                    )

                except PermissionError:
                    st.error(
                        "Não foi possível salvar. Feche a planilha "
                        "enderecos_logisticos.xlsx no Excel e tente novamente."
                    )

                else:
                    st.success(
                        f"O endereço {codigo_para_liberar} "
                        "foi liberado com sucesso!"
                    )

                    st.rerun()

with pagina_historico:
    st.subheader("Histórico de movimentações")

    st.write(
        "Consulte os cadastros, entradas, saídas "
        "e liberações realizados no sistema."
    )

    historico = carregar_historico()

    if historico.empty:
        st.info(
            "Ainda não existem movimentações registradas."
        )

    else:
        total_movimentacoes = len(historico)

        total_cadastros = len(
            historico[historico["Tipo"] == "Cadastro"]
        )

        total_entradas = len(
            historico[historico["Tipo"] == "Entrada"]
        )

        total_saidas = len(
            historico[historico["Tipo"] == "Saída"]
        )

        total_liberacoes = len(
            historico[historico["Tipo"] == "Liberação"]
        )

        coluna_h1, coluna_h2, coluna_h3, coluna_h4, coluna_h5 = (
            st.columns(5)
        )

        coluna_h1.metric(
            "Movimentações",
            total_movimentacoes,
        )

        coluna_h2.metric(
            "Cadastros",
            total_cadastros,
        )

        coluna_h3.metric(
            "Entradas",
            total_entradas,
        )

        coluna_h4.metric(
            "Saídas",
            total_saidas,
        )

        coluna_h5.metric(
            "Liberações",
            total_liberacoes,
        )

        st.divider()

        st.write("### Filtros do histórico")

        filtro_coluna1, filtro_coluna2 = st.columns(2)

        tipos_existentes = (
            historico["Tipo"]
            .dropna()
            .astype(str)
            .unique()
            .tolist()
        )

        tipos_existentes = sorted(tipos_existentes)

        with filtro_coluna1:
            tipo_historico = st.selectbox(
                "Tipo de movimentação",
                ["Todos"] + tipos_existentes,
                key="filtro_tipo_historico",
            )

        with filtro_coluna2:
            pesquisa_historico = st.text_input(
                "Pesquisar endereço ou produto",
                placeholder="Exemplo: B01-R01-BA02 ou Monitor",
                key="pesquisa_historico",
            )

        historico_filtrado = historico.copy()

        if tipo_historico != "Todos":
            historico_filtrado = historico_filtrado[
                historico_filtrado["Tipo"]
                == tipo_historico
            ]

        if pesquisa_historico:
            texto_pesquisado = pesquisa_historico.strip()

            filtro_endereco = (
                historico_filtrado["Endereço"]
                .astype(str)
                .str.contains(
                    texto_pesquisado,
                    case=False,
                    na=False,
                )
            )

            filtro_produto = (
                historico_filtrado["Produto"]
                .astype(str)
                .str.contains(
                    texto_pesquisado,
                    case=False,
                    na=False,
                )
            )

            historico_filtrado = historico_filtrado[
                filtro_endereco | filtro_produto
            ]

        historico_filtrado = historico_filtrado.sort_values(
            by="Data e hora",
            ascending=False,
        )

        tabela_historico = historico_filtrado.copy()

        tabela_historico["Data e hora"] = (
            tabela_historico["Data e hora"]
            .dt.strftime("%d/%m/%Y %H:%M:%S")
            .fillna("")
        )

        st.write(
            f"**Registros encontrados:** "
            f"{len(tabela_historico)}"
        )

        st.dataframe(
            tabela_historico,
            width="stretch",
            hide_index=True,
        )


with pagina_dashboard:
    st.subheader("Dashboard Logístico")

    st.write(
        "Acompanhe a ocupação dos endereços, "
        "as quantidades armazenadas e as movimentações."
    )

    total_dashboard = len(enderecos)

    ocupados_dashboard = len(
        enderecos[enderecos["status"] == "Ocupado"]
    )

    disponiveis_dashboard = len(
        enderecos[enderecos["status"] == "Disponível"]
    )

    quantidade_total = int(
        enderecos["quantidade"].fillna(0).sum()
    )

    if total_dashboard > 0:
        taxa_dashboard = (
            ocupados_dashboard / total_dashboard * 100
        )
    else:
        taxa_dashboard = 0

    coluna_d1, coluna_d2, coluna_d3, coluna_d4, coluna_d5 = (
        st.columns(5)
    )

    coluna_d1.metric(
        "Total de endereços",
        total_dashboard,
    )

    coluna_d2.metric(
        "Disponíveis",
        disponiveis_dashboard,
    )

    coluna_d3.metric(
        "Ocupados",
        ocupados_dashboard,
    )

    coluna_d4.metric(
        "Taxa de ocupação",
        f"{taxa_dashboard:.1f}%",
    )

    coluna_d5.metric(
        "Unidades armazenadas",
        quantidade_total,
    )

    st.divider()

    coluna_grafico1, coluna_grafico2 = st.columns(2)

    with coluna_grafico1:
        st.write("### Endereços ocupados por buffer")

        ocupacao_buffer = (
            enderecos[
                enderecos["status"] == "Ocupado"
            ]
            .groupby("buffer")
            .size()
            .reset_index(name="Endereços ocupados")
        )

        todos_buffers = pd.DataFrame(
            {
                "buffer": sorted(
                    enderecos["buffer"].unique()
                )
            }
        )

        ocupacao_buffer = todos_buffers.merge(
            ocupacao_buffer,
            on="buffer",
            how="left",
        )

        ocupacao_buffer[
            "Endereços ocupados"
        ] = ocupacao_buffer[
            "Endereços ocupados"
        ].fillna(0)

        ocupacao_buffer = ocupacao_buffer.set_index(
            "buffer"
        )

        st.bar_chart(
            ocupacao_buffer,
            y="Endereços ocupados",
        )

    with coluna_grafico2:
        st.write("### Quantidade armazenada por buffer")

        quantidade_buffer = (
            enderecos
            .groupby("buffer")["quantidade"]
            .sum()
            .reset_index()
        )

        quantidade_buffer = quantidade_buffer.rename(
            columns={
                "quantidade": "Quantidade armazenada",
            }
        )

        quantidade_buffer = quantidade_buffer.set_index(
            "buffer"
        )

        st.bar_chart(
            quantidade_buffer,
            y="Quantidade armazenada",
        )

    st.divider()

    st.write("### Ocupação por rua")

    ocupacao_rua = (
        enderecos[
            enderecos["status"] == "Ocupado"
        ]
        .groupby("rua")
        .size()
        .reset_index(name="Endereços ocupados")
    )

    todas_ruas = pd.DataFrame(
        {
            "rua": sorted(
                enderecos["rua"].unique()
            )
        }
    )

    ocupacao_rua = todas_ruas.merge(
        ocupacao_rua,
        on="rua",
        how="left",
    )

    ocupacao_rua[
        "Endereços ocupados"
    ] = ocupacao_rua[
        "Endereços ocupados"
    ].fillna(0)

    ocupacao_rua = ocupacao_rua.set_index("rua")

    st.bar_chart(
        ocupacao_rua,
        y="Endereços ocupados",
    )

    st.divider()

    coluna_produtos, coluna_movimentacoes = st.columns(2)

    with coluna_produtos:
        st.write("### Produtos com maiores quantidades")

        produtos_armazenados = enderecos[
            enderecos["status"] == "Ocupado"
        ].copy()

        if produtos_armazenados.empty:
            st.info(
                "Não existem produtos armazenados."
            )

        else:
            ranking_produtos = (
                produtos_armazenados
                .groupby("produto")["quantidade"]
                .sum()
                .reset_index()
                .sort_values(
                    by="quantidade",
                    ascending=False,
                )
                .head(10)
            )

            ranking_produtos = ranking_produtos.rename(
                columns={
                    "produto": "Produto",
                    "quantidade": "Quantidade",
                }
            )

            ranking_produtos = ranking_produtos.set_index(
                "Produto"
            )

            st.bar_chart(
                ranking_produtos,
                y="Quantidade",
            )

    with coluna_movimentacoes:
        st.write("### Movimentações por tipo")

        historico_dashboard = carregar_historico()

        if historico_dashboard.empty:
            st.info(
                "Ainda não existem movimentações "
                "registradas."
            )

        else:
            movimentacoes_tipo = (
                historico_dashboard
                .groupby("Tipo")
                .size()
                .reset_index(name="Movimentações")
            )

            movimentacoes_tipo = (
                movimentacoes_tipo.set_index("Tipo")
            )

            st.bar_chart(
                movimentacoes_tipo,
                y="Movimentações",
            )

    st.divider()

    st.write("### Resumo da ocupação por buffer")

    resumo_buffer = (
        enderecos
        .groupby(["buffer", "status"])
        .size()
        .unstack(fill_value=0)
        .reset_index()
    )

    resumo_buffer = resumo_buffer.rename(
        columns={
            "buffer": "Buffer",
            "Disponível": "Disponíveis",
            "Ocupado": "Ocupados",
        }
    )

    if "Disponíveis" not in resumo_buffer.columns:
        resumo_buffer["Disponíveis"] = 0

    if "Ocupados" not in resumo_buffer.columns:
        resumo_buffer["Ocupados"] = 0

    resumo_buffer["Total"] = (
        resumo_buffer["Disponíveis"]
        + resumo_buffer["Ocupados"]
    )

    resumo_buffer["Taxa de ocupação"] = (
        resumo_buffer["Ocupados"]
        / resumo_buffer["Total"]
        * 100
    ).round(1)

    resumo_buffer["Taxa de ocupação"] = (
        resumo_buffer["Taxa de ocupação"]
        .astype(str)
        + "%"
    )

    st.dataframe(
        resumo_buffer[
            [
                "Buffer",
                "Disponíveis",
                "Ocupados",
                "Total",
                "Taxa de ocupação",
            ]
        ],
        width="stretch",
        hide_index=True,
    )

with pagina_relatorios:
    st.subheader("Relatórios e exportações")

    st.write(
        "Gere e baixe relatórios dos endereços "
        "e das movimentações do sistema."
    )

    st.divider()

    tabela_enderecos_exportacao = enderecos.rename(
        columns={
            "codigo": "Código",
            "buffer": "Buffer",
            "rua": "Rua",
            "canalizacao": "Canalização",
            "status": "Status",
            "produto": "Produto",
            "quantidade": "Quantidade",
        }
    )

    tabela_ocupados = tabela_enderecos_exportacao[
        tabela_enderecos_exportacao["Status"]
        == "Ocupado"
    ].copy()

    tabela_disponiveis = tabela_enderecos_exportacao[
        tabela_enderecos_exportacao["Status"]
        == "Disponível"
    ].copy()

    historico_relatorio = carregar_historico()

    total_relatorio = len(
        tabela_enderecos_exportacao
    )

    ocupados_relatorio = len(tabela_ocupados)
    disponiveis_relatorio = len(tabela_disponiveis)
    movimentacoes_relatorio = len(historico_relatorio)

    coluna_r1, coluna_r2, coluna_r3, coluna_r4 = (
        st.columns(4)
    )

    coluna_r1.metric(
        "Todos os endereços",
        total_relatorio,
    )

    coluna_r2.metric(
        "Endereços ocupados",
        ocupados_relatorio,
    )

    coluna_r3.metric(
        "Endereços disponíveis",
        disponiveis_relatorio,
    )

    coluna_r4.metric(
        "Movimentações",
        movimentacoes_relatorio,
    )

    st.divider()

    st.write("### Relatórios de endereçamento")

    coluna_download1, coluna_download2, coluna_download3 = (
        st.columns(3)
    )

    with coluna_download1:
        arquivo_todos = criar_excel_para_download(
            tabela_enderecos_exportacao,
            "Todos os endereços",
        )

        st.download_button(
            label="Baixar todos os endereços",
            data=arquivo_todos,
            file_name="relatorio_todos_enderecos.xlsx",
            mime=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
            type="primary",
        )

    with coluna_download2:
        arquivo_ocupados = criar_excel_para_download(
            tabela_ocupados,
            "Endereços ocupados",
        )

        st.download_button(
            label="Baixar endereços ocupados",
            data=arquivo_ocupados,
            file_name="relatorio_enderecos_ocupados.xlsx",
            mime=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
        )

    with coluna_download3:
        arquivo_disponiveis = criar_excel_para_download(
            tabela_disponiveis,
            "Endereços disponíveis",
        )

        st.download_button(
            label="Baixar endereços disponíveis",
            data=arquivo_disponiveis,
            file_name="relatorio_enderecos_disponiveis.xlsx",
            mime=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
        )

    st.divider()

    st.write("### Histórico de movimentações")

    if historico_relatorio.empty:
        st.info(
            "Ainda não existem movimentações "
            "para exportar."
        )

    else:
        historico_para_download = (
            historico_relatorio.copy()
        )

        historico_para_download[
            "Data e hora"
        ] = historico_para_download[
            "Data e hora"
        ].dt.strftime(
            "%d/%m/%Y %H:%M:%S"
        )

        arquivo_historico = criar_excel_para_download(
            historico_para_download,
            "Histórico",
        )

        st.download_button(
            label="Baixar histórico completo",
            data=arquivo_historico,
            file_name=(
                "relatorio_historico_movimentacoes.xlsx"
            ),
            mime=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
            type="primary",
        )

        st.dataframe(
            historico_para_download,
            width="stretch",
            hide_index=True,
        )    